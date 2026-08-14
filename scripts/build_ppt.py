# -*- coding: utf-8 -*-
"""生成《智服方舟——多智能体客服自主闭环平台》初赛方案 PPT。"""
import math
import os

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from pptkit import (new_prs, blank_slide, header, add_text, rich_text, para, rect, label,
                    arrow, add_table, section_badge, circle,
                    NAVY, NAVY_DEEP, TEAL, TEAL_DARK, BLUE, LBLUE, WHITE, OFFWHITE,
                    INK, GRAY, LGRAY, AMBER, RED, GREEN, FONT, W, H, MARGIN, set_run,
                    text_width_in)

CYAN = RGBColor(0x5E, 0xD0, 0xCE)
GRAY_LIGHT = RGBColor(0x9A, 0xAB, 0xBE)

TOTAL = 19


def s01_cover(prs):
    slide = blank_slide(prs, bg=NAVY_DEEP)
    # 装饰
    circle(slide, Inches(11.9), Inches(1.15), Inches(2.7), line=RGBColor(0x1B, 0x3A, 0x5F), line_w=1.5)
    circle(slide, Inches(12.7), Inches(0.45), Inches(0.9), fill=TEAL)
    circle(slide, Inches(0.9), Inches(6.6), Inches(0.5), fill=BLUE)
    circle(slide, Inches(1.6), Inches(6.95), Inches(0.24), fill=AMBER)
    # 主体
    rect(slide, Inches(1.1), Inches(2.0), Inches(0.14), Inches(0.95), fill=TEAL)
    add_text(slide, Inches(1.42), Inches(1.9), Inches(9.5), Inches(1.1),
             [[("智服方舟", 48, True, WHITE)]])
    add_text(slide, Inches(1.44), Inches(2.98), Inches(9.5), Inches(0.7),
             [[("多智能体客服自主闭环平台", 26, True, CYAN)]])
    add_text(slide, Inches(1.46), Inches(3.72), Inches(9.5), Inches(0.4),
             [[("Multi-Agent Customer Service Autonomous Closed-Loop Platform", 12, False, GRAY_LIGHT)]])
    rect(slide, Inches(1.1), Inches(4.32), Inches(4.2), Pt(1.5), fill=RGBColor(0x2A, 0x4A, 0x6E))
    add_text(slide, Inches(1.1), Inches(4.5), Inches(9.5), Inches(0.4),
             [[("方向二 · 智能客服自主闭环", 15, True, WHITE),
               ("　｜　初赛：方向与方案设计", 13, False, GRAY_LIGHT)]])
    # 技术标签
    tags = ["AgentTeams 协同底座", "Leader-Worker", "Skill", "MCP", "RAG", "可观测"]
    x = Inches(1.1)
    for t in tags:
        w = Inches(0.32 + 0.18 * len(t))
        label(slide, x, Inches(5.35), w, Inches(0.42), t, NAVY, color=CYAN, size=10.5,
              line=RGBColor(0x2A, 0x4A, 0x6E))
        x += w + Inches(0.18)
    add_text(slide, Inches(1.1), Inches(6.35), Inches(10), Inches(0.4),
             [[("让客服从“单点问答”走向“自主闭环”：接入 → 分诊 → 方案 → 执行 → 核验 → 复盘", 12.5, False, WHITE)]])
    add_text(slide, Inches(1.1), Inches(6.85), Inches(10), Inches(0.35),
             [[("参赛团队：____________　　日期：2026-08", 10, False, GRAY_LIGHT)]])
    return slide


def s02_toc(prs):
    slide = blank_slide(prs)
    header(slide, "目录", page=2, total=TOTAL)
    items = [
        ("01", "背景与业务价值", "多渠道客服的五大痛点与量化目标"),
        ("02", "总体方案架构", "接入 / 协同 / 能力 / 数据 / 治理五层"),
        ("03", "多 Agent 协同设计", "Leader-Worker 角色编排与 AgentTeams 映射"),
        ("04", "端到端闭环流程", "八环节闭环：输入 → 拆解 → 上下文 → 工具 → 验证 → 证据 → 审批 → 沉淀"),
        ("05", "Skill 与 MCP 工具集成", "7 个核心 Skill 九要素清单与工具契约"),
        ("06", "可观测与 RAG 增强", "Trace/Log/Metrics 与四类上下文能力"),
        ("07", "安全与异常分支", "审批 / 回滚 / 审计 / 容错降级"),
        ("08", "可行性与落地计划", "初赛 → 复赛 → 决赛 → 生产四阶段"),
        ("09", "项目进展与总结", "当前进展、开放开源与下一步"),
    ]
    x0, y0 = Inches(0.55), Inches(1.72)
    cw, ch = Inches(3.94), Inches(1.52)
    gx, gy = Inches(0.24), Inches(0.22)
    for i, (num, title, desc) in enumerate(items):
        r, c = divmod(i, 3)
        x = x0 + c * (cw + gx)
        y = y0 + r * (ch + gy)
        rect(slide, x, y, cw, ch, fill=OFFWHITE, line=LGRAY, line_w=1.0,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        label(slide, x + Inches(0.18), y + Inches(0.18), Inches(0.62), Inches(0.5),
              num, TEAL, size=15, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
        add_text(slide, x + Inches(0.98), y + Inches(0.2), cw - Inches(1.15), Inches(0.5),
                 [[(title, 14.5, True, NAVY)]])
        add_text(slide, x + Inches(0.2), y + Inches(0.88), cw - Inches(0.38), Inches(0.55),
                 [[(desc, 9.5, False, GRAY)]])
    return slide


def pain_row(slide, y, title, desc, idx):
    rect(slide, Inches(5.25), y, Inches(7.5), Inches(0.74), fill=OFFWHITE, line=LGRAY,
         line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    circle(slide, Inches(5.5) + Inches(0.19), y + Inches(0.37), Inches(0.38), fill=TEAL)
    add_text(slide, Inches(5.48), y + Inches(0.24), Inches(0.42), Inches(0.3),
             [[(str(idx), 12, True, WHITE)]], align=PP_ALIGN.CENTER)
    add_text(slide, Inches(5.98), y + Inches(0.08), Inches(6.7), Inches(0.3),
             [[(title, 12.5, True, NAVY)]])
    add_text(slide, Inches(5.98), y + Inches(0.4), Inches(6.7), Inches(0.3),
             [[(desc, 9.5, False, GRAY)]])


def s03_problem(prs):
    slide = blank_slide(prs)
    header(slide, "背景与业务价值", subtitle="多渠道客服场景：重复、分诊慢、跨系统执行慢、缺核验、难沉淀", page=3, total=TOTAL)
    # 左：场景
    rect(slide, MARGIN, Inches(1.55), Inches(4.35), Inches(4.35), fill=OFFWHITE, line=LGRAY,
         line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    add_text(slide, Inches(0.85), Inches(1.75), Inches(3.8), Inches(0.4),
             [[("业务场景：多渠道会话聚合", 15, True, TEAL_DARK)]])
    channels = [
        ("邮件渠道", "Mail · 工单导入"),
        ("在线客服", "Web IM · 实时会话"),
        ("电话转写", "ASR 通话文本"),
        ("社媒评论", "微博/小红书/微信"),
        ("企业工单系统", "既有 Ticketing"),
    ]
    y = Inches(2.32)
    for name, sub in channels:
        rect(slide, Inches(0.85), y, Inches(3.75), Inches(0.56), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
        rect(slide, Inches(1.05), y + Inches(0.18), Inches(0.14), Inches(0.2), fill=TEAL)
        add_text(slide, Inches(1.35), y + Inches(0.07), Inches(2.2), Inches(0.3),
                 [[(name, 11.5, True, NAVY)]])
        add_text(slide, Inches(1.35), y + Inches(0.3), Inches(3.1), Inches(0.25),
                 [[(sub, 8.5, False, GRAY)]])
        y += Inches(0.68)
    add_text(slide, Inches(0.85), y + Inches(0.04), Inches(3.85), Inches(0.75),
             [[("同一用户多渠道重复发起咨询", 10, True, AMBER)],
              [("→ 会话指纹去重、工单合并", 10, False, GRAY)]], )
    # 右：痛点
    add_text(slide, Inches(5.25), Inches(1.55), Inches(4), Inches(0.4),
             [[("五大业务痛点", 15, True, NAVY)]])
    pains = [
        ("渠道割裂，重复工单率高", "消息散落多渠道，无统一聚合与去重"),
        ("人工分诊，响应慢且标准不一", "意图依赖经验，紧急/高危识别不及时"),
        ("跨系统执行慢", "退款、换货、账户变更需多系统人工操作"),
        ("执行结果缺核验", "无自动核验，满意度确认滞后"),
        ("经验难沉淀", "疑难案例复盘靠手工，知识库更新慢"),
    ]
    y = Inches(2.05)
    for i, (t, d) in enumerate(pains, 1):
        pain_row(slide, y, t, d, i)
        y += Inches(0.82)
    # 价值目标条
    rect(slide, MARGIN, Inches(6.25), W - 2 * MARGIN, Inches(0.72), fill=NAVY,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    label(slide, Inches(0.75), Inches(6.36), Inches(1.5), Inches(0.5), "价值目标", TEAL, size=11)
    metrics = ["工单处理时效 ↓70%", "重复工单率 ↓60%", "一次解决率 ↑40%", "人工介入率 ≤15%（仅高风险/异常）"]
    x = Inches(2.55)
    for m in metrics:
        rect(slide, x, Inches(6.33), Inches(2.42), Inches(0.56), fill=RGBColor(0x1B, 0x3A, 0x5F),
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
        add_text(slide, x + Inches(0.08), Inches(6.41), Inches(2.28), Inches(0.42),
                 [[(m, 10, True, WHITE)]], align=PP_ALIGN.CENTER)
        x += Inches(2.56)
    return slide


def band(slide, y, h, title, content_cb, fill=OFFWHITE, title_color=NAVY):
    rect(slide, MARGIN, y, W - 2 * MARGIN, h, fill=fill, line=LGRAY, line_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    rect(slide, MARGIN + Inches(0.12), y + Inches(0.12), Inches(0.09), h - Inches(0.24), fill=title_color)
    add_text(slide, MARGIN + Inches(0.32), y + Inches(0.12), Inches(1.15), h - Inches(0.24),
             [[(title, 11.5, True, NAVY)]], anchor=MSO_ANCHOR.MIDDLE)
    content_cb(slide, MARGIN + Inches(1.5), y)


def chip_row(slide, x, y, h, items, w, gap=0.14, fill=WHITE, color=NAVY, size=10, bold=True):
    cx = x
    for t in items:
        label(slide, cx, y, w, h, t, fill, color=color, size=size, bold=bold, line=LGRAY)
        cx += w + gap
    return cx


def s04_arch(prs):
    slide = blank_slide(prs)
    header(slide, "总体方案架构", subtitle="以 AgentTeams 为协同底座：接入 / 协同 / 能力 / 数据 / 治理五层", page=4, total=TOTAL)

    def band1(s, x, y):
        chip_row(s, x, y + Inches(0.11), Inches(0.4),
                 ["邮件", "在线客服", "电话转写", "社媒评论", "工单系统"], Inches(1.9), fill=WHITE)

    def band2(s, x, y):
        label(s, x, y + Inches(0.1), Inches(2.85), Inches(0.6), "主控 Agent（Leader）\n编排 · 状态追踪", NAVY, size=9.5)
        workers = ["聚合", "分诊", "方案", "执行", "核验", "复盘"]
        ww, wh = Inches(1.5), Inches(0.42)
        for i, name in enumerate(workers):
            col, row = divmod(i, 3)
            label(s, x + Inches(3.05) + col * (ww + Inches(0.12)), y + Inches(0.12) + row * (wh + Inches(0.12)),
                  ww, wh, name, LBLUE, color=NAVY, size=10.5)
        add_text(s, x + Inches(8.1), y + Inches(0.1), Inches(2.6), Inches(1.35),
                 [[("Matrix 房间传递任务与中间结论", 8.5, False, GRAY)],
                  [("心跳监控 Worker 运行状态", 8.5, False, GRAY)],
                  [("RBAC 权限与人工介入闸口", 8.5, False, GRAY)],
                  [("Worker 模板批量部署", 8.5, False, GRAY)]],
                 )

    def band3(s, x, y):
        label(s, x, y + Inches(0.1), Inches(4.5), Inches(0.6),
              "Skill 能力抽象层\n7 个核心 Skill（接入/认知/决策/执行/核验/沉淀）", LBLUE, color=NAVY, size=9)
        arrow(s, x + Inches(4.72), y + Inches(0.4), x + Inches(5.2), y + Inches(0.4), color=TEAL)
        label(s, x + Inches(5.4), y + Inches(0.1), Inches(5.0), Inches(0.6),
              "MCP 工具连接层\n工单 · 支付 · ERP · CRM · 知识库 · 短信/IM", TEAL, size=9)

    def band4(s, x, y):
        chip_row(s, x, y + Inches(0.14), Inches(0.42),
                 ["知识库 RAG（向量库）", "工单 / 案例库", "Agent 记忆（Redis）", "观测存储（SLS · OTLP）"],
                 Inches(2.35), fill=WHITE)

    def band5(s, x, y):
        chip_row(s, x, y + Inches(0.11), Inches(0.4),
                 ["AgentLoop 全链路观测", "审计日志", "人工审批台", "模型管理（多供应商）", "告警与离线评估"],
                 Inches(1.85), fill=OFFWHITE)

    band(slide, Inches(1.5), Inches(0.62), "接入层", band1, fill=OFFWHITE)
    band(slide, Inches(2.22), Inches(1.62), "协同层\nAgentTeams", band2, fill=WHITE)
    band(slide, Inches(3.94), Inches(0.9), "能力层", band3, fill=OFFWHITE)
    band(slide, Inches(4.94), Inches(0.72), "数据层", band4, fill=WHITE)
    band(slide, Inches(5.76), Inches(0.62), "治理·\n可观测", band5, fill=OFFWHITE)
    add_text(slide, MARGIN, Inches(6.55), W - 2 * MARGIN, Inches(0.35),
             [[("复用 AgentTeams 平台能力：多模型供应商接入、MCP 服务管理、Worker 模板、监控仪表盘、团队权限隔离",
                 10, False, GRAY)]])
    return slide


def s05_roles(prs):
    slide = blank_slide(prs)
    header(slide, "多 Agent 协同设计：Leader-Worker 角色编排",
           subtitle="1 个 Leader + 6 个职能 Worker，分工明确、边界清晰", page=5, total=TOTAL)
    # Leader
    rect(slide, Inches(4.53), Inches(1.5), Inches(4.27), Inches(0.95), fill=NAVY,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    add_text(slide, Inches(4.73), Inches(1.58), Inches(3.9), Inches(0.4),
             [[("主控 Agent · Leader", 14.5, True, WHITE)]])
    add_text(slide, Inches(4.73), Inches(2.02), Inches(4.0), Inches(0.4),
             [[("任务拆解 / 编排调度 / 上下文路由 / 状态追踪 / 升级决策", 9, False, CYAN)]])
    # 三个职能组
    groups = [
        (Inches(0.55), TEAL, "接入与认知", [
            ("聚合 Agent · Intake", "多渠道接入 / 归一化 / 去重建档"),
            ("分诊 Agent · Triage", "意图识别 / 分级 / 信息补全"),
        ]),
        (Inches(4.69), BLUE, "决策与执行", [
            ("方案 Agent · Planner", "RAG 检索 / 生成合规方案"),
            ("执行 Agent · Executor", "Skill+MCP 执行 / 幂等 / 回滚"),
        ]),
        (Inches(8.83), AMBER, "验证与进化", [
            ("核验 Agent · Verifier", "结果核验 / 满意度确认"),
            ("复盘 Agent · Reviewer", "根因分析 / 知识沉淀"),
        ]),
    ]
    for x, color, gname, agents in groups:
        rect(slide, x, Inches(2.78), Inches(3.95), Inches(2.62), fill=OFFWHITE, line=LGRAY,
             line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        rect(slide, x, Inches(2.78), Inches(3.95), Inches(0.44), fill=color,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.35)
        add_text(slide, x, Inches(2.83), Inches(3.95), Inches(0.34),
                 [[(gname, 12, True, WHITE)]], align=PP_ALIGN.CENTER)
        for j, (name, duty) in enumerate(agents):
            ay = Inches(3.36) + j * Inches(1.0)
            rect(slide, x + Inches(0.18), ay, Inches(3.59), Inches(0.88), fill=WHITE, line=LGRAY,
                 line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
            rect(slide, x + Inches(0.18), ay, Inches(0.07), Inches(0.88), fill=color)
            add_text(slide, x + Inches(0.36), ay + Inches(0.1), Inches(3.3), Inches(0.35),
                     [[(name, 11.5, True, NAVY)]])
            add_text(slide, x + Inches(0.36), ay + Inches(0.46), Inches(3.3), Inches(0.36),
                     [[(duty, 8.5, False, GRAY)]])
        arrow(slide, Inches(6.665), Inches(2.45), x + Inches(1.975), Inches(2.78), color=GRAY, weight=1.5)
    # 协同三原则
    rect(slide, MARGIN, Inches(5.62), W - 2 * MARGIN, Inches(0.68), fill=RGBColor(0xE6, 0xF5, 0xF4),
         line=TEAL, line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    add_text(slide, MARGIN + Inches(0.2), Inches(5.7), W - 2 * MARGIN - Inches(0.4), Inches(0.55),
             [[("协同三原则：", 11, True, TEAL_DARK),
               ("① 单一领导权——一单一主控，避免双头指挥　　", 10.5, False, INK),
               ("② 最小授权——Worker 仅持本职权限，跨权限须经主控路由　　", 10.5, False, INK),
               ("③ 人工兜底——高风险 / 低置信 / 异常升级 Human-in-the-loop", 10.5, False, INK)]])
    add_text(slide, MARGIN, Inches(6.5), W - 2 * MARGIN, Inches(0.35),
             [[("身份属性、能力边界与协同关系完整清单见《附录 A：Agent Identity 清单》", 10, False, GRAY)]])
    return slide


def s06_identity(prs):
    slide = blank_slide(prs)
    header(slide, "Agent Identity 清单（摘要）",
           subtitle="对应《参赛手册-附录A》：身份属性 / 能力边界 / 协同关系 / 权限约束", page=6, total=TOTAL)
    rows = [
        ["Agent", "身份定位", "能力边界（能做 / 不做）", "协同关系"],
        ["主控 Agent（Leader）", "值班长：任务拆解、编排调度、全局状态追踪",
         "做：编排、路由、仲裁、升级决策、结果汇总；不做：直接操作业务系统",
         "与全部 Worker 双向协作；向人工发起审批"],
        ["聚合 Agent（Intake）", "前台接待：多渠道接入、归一化、去重、建档",
         "做：解析消息、指纹去重、生成标准工单；不做：意图判断与业务处理",
         "上游多渠道；下游分诊 Agent"],
        ["分诊 Agent（Triage）", "分诊护士：意图识别、分类分级、补全信息",
         "做：意图/槽位提取、紧急度分级；不做：生成处理动作",
         "上游聚合；下游方案 Agent；低置信转人工"],
        ["方案 Agent（Planner）", "政策顾问：RAG 检索、生成合规方案",
         "做：检索 SOP/案例、出方案并标注证据与风险；不做：直接执行",
         "上游分诊；下游执行 Agent；高风险发起审批"],
        ["执行 Agent（Executor）", "柜面操作员：Skill/MCP 执行、幂等、回滚",
         "做：按批准方案调用系统、重试回滚；不做：擅自改方案、超授权",
         "上游方案（须带审批凭证）；下游核验 Agent"],
        ["核验 Agent（Verifier）", "质检员：结果核验、满意度确认、证据包",
         "做：比对证据、触发满意度确认、出核验报告；不做：重新执行",
         "上游执行；输出核验报告交主控归档"],
        ["复盘 Agent（Reviewer）", "分析师：根因分析、知识沉淀",
         "做：疑难案例复盘、知识库与 Skill 策略建议；不做：处理实时工单",
         "上游核验/主控；知识写入需人工确认"],
    ]
    add_table(slide, MARGIN, Inches(1.5), W - 2 * MARGIN, rows,
              [1.35, 2.9, 4.6, 3.38], row_height=0.6, font_size=9)
    add_text(slide, MARGIN, Inches(6.55), W - 2 * MARGIN, Inches(0.4),
             [[("权限约束：主控仅调度不写业务；执行 Agent 持最小业务 Token；所有操作留痕审计；敏感数据先脱敏后入库",
                 10, False, GRAY)]])
    return slide


def step_box(slide, x, y, w, h, idx, title, agent, desc, color=TEAL):
    rect(slide, x, y, w, h, fill=WHITE, line=LGRAY, line_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.09)
    rect(slide, x, y, Inches(0.5), h, fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    add_text(slide, x, y + Inches(0.04), Inches(0.5), Inches(0.35),
             [[(str(idx), 13, True, WHITE)]], align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.64), y + Inches(0.1), w - Inches(0.8), Inches(0.35),
             [[(title, 12.5, True, NAVY)]])
    add_text(slide, x + Inches(0.64), y + Inches(0.5), w - Inches(0.8), Inches(0.3),
             [[(agent, 9.5, True, color if color in (TEAL, BLUE, AMBER) else TEAL_DARK)]])
    add_text(slide, x + Inches(0.64), y + Inches(0.85), w - Inches(0.8), Inches(0.55),
             [[(desc, 8.5, False, GRAY)]])


def s07_loop(prs):
    slide = blank_slide(prs)
    header(slide, "端到端闭环流程：八环节自主闭环",
           subtitle="对应赛题 1.3：输入 / 拆解 / 上下文 / 工具 / 验证 / 证据 / 审批回滚 / 经验沉淀", page=7, total=TOTAL)
    steps = [
        ("任务输入", "聚合 Agent", "多渠道消息/工单/告警\n归一化与去重"),
        ("任务拆解", "主控 Agent", "目标理解、任务分解\n与优先级编排"),
        ("上下文传递", "AgentTeams", "Matrix 房间 +\nTicket Context 共享状态"),
        ("工具调用", "执行 Agent", "Skill + MCP\n调用企业系统"),
        ("结果验证", "核验 Agent", "证据比对\n满意度确认"),
        ("证据沉淀", "主控 Agent", "Trace / Log\n证据包归档"),
        ("审批与回滚", "人工 + 主控", "高风险人工审批\n幂等执行与回滚"),
        ("经验沉淀", "复盘 Agent", "知识库更新\nSkill 策略优化"),
    ]
    xs = [Inches(0.55), Inches(3.28), Inches(6.01), Inches(8.74)]
    y1, y2 = Inches(1.62), Inches(3.82)
    for i, (t, a, d) in enumerate(steps):
        col = i % 4
        row = 0 if i < 4 else 1
        y = y1 if row == 0 else y2
        step_box(slide, xs[col], y, Inches(2.55), Inches(1.55), i + 1, t, a, d)
    # 箭头：第一行右向
    for c in range(3):
        arrow(slide, xs[c] + Inches(2.55), y1 + Inches(0.78), xs[c + 1], y1 + Inches(0.78), color=TEAL, weight=2)
    # 第二行左向
    for c in range(3):
        arrow(slide, xs[3 - c], y2 + Inches(0.78), xs[2 - c] + Inches(2.55), y2 + Inches(0.78), color=TEAL, weight=2)
    # 连接：4→8 与 5→1
    arrow(slide, xs[3] + Inches(2.55), y1 + Inches(1.55), xs[3] + Inches(2.55), y2, color=TEAL, weight=2)
    arrow(slide, xs[0], y2, xs[0], y1 + Inches(1.55), color=AMBER, weight=2)
    # 中间徽章
    rect(slide, Inches(4.6), Inches(3.02), Inches(4.1), Inches(0.62), fill=NAVY,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    add_text(slide, Inches(4.7), Inches(3.09), Inches(3.9), Inches(0.5),
             [[("AgentTeams 编排 · 工单状态机 · 全链路追踪", 11, True, WHITE)]],
             align=PP_ALIGN.CENTER)
    # 底部：闭环要求 → 设计落点
    rect(slide, MARGIN, Inches(5.72), W - 2 * MARGIN, Inches(1.0), fill=OFFWHITE, line=LGRAY,
         line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    add_text(slide, MARGIN + Inches(0.2), Inches(5.82), Inches(2.2), Inches(0.4),
             [[("闭环要求 → 设计落点", 12, True, NAVY)]])
    maps = [
        ("任务输入/拆解", "聚合 + 主控 Agent"),
        ("上下文传递/工具调用", "Matrix + Ticket Context + Skill/MCP"),
        ("结果验证/证据沉淀", "核验 Agent + Trace/Log/证据包"),
        ("审批回滚/经验沉淀", "人工审批 + 幂等回滚 + 复盘 Agent"),
    ]
    x = MARGIN + Inches(2.5)
    for t, v in maps:
        rect(slide, x, Inches(5.85), Inches(2.42), Inches(0.78), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        add_text(slide, x + Inches(0.12), Inches(5.92), Inches(2.2), Inches(0.3),
                 [[(t, 9.5, True, TEAL_DARK)]])
        add_text(slide, x + Inches(0.12), Inches(6.24), Inches(2.2), Inches(0.4),
                 [[(v, 8.5, False, GRAY)]])
        x += Inches(2.56)
    return slide


def s08_context(prs):
    slide = blank_slide(prs)
    header(slide, "上下文传递与状态追踪",
           subtitle="Ticket Context 作为单一事实源（SSOT），由 Leader 维护状态机", page=8, total=TOTAL)
    # 左：Ticket Context
    rect(slide, MARGIN, Inches(1.5), Inches(6.0), Inches(5.15), fill=OFFWHITE, line=LGRAY,
         line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(slide, MARGIN, Inches(1.5), Inches(6.0), Inches(0.5), fill=NAVY,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
    add_text(slide, MARGIN + Inches(0.2), Inches(1.56), Inches(5.6), Inches(0.4),
             [[("工单上下文 Ticket Context（共享状态）", 13, True, WHITE)]])
    fields = [
        ("ticket_id", "唯一标识 · 关联 TraceID"),
        ("原始会话", "多渠道原文（脱敏后）· 会话指纹去重记录"),
        ("意图与分级", "intent / 紧急度 / 风险等级 / 情绪"),
        ("方案与审批", "处理方案 · 证据引用 · 审批状态"),
        ("执行证据", "工具调用结果 · 凭证 · 幂等键"),
        ("核验结果", "核验报告 · 满意度反馈"),
        ("复盘标签", "案例标签 · 知识库关联"),
    ]
    y = Inches(2.14)
    for k, v in fields:
        rect(slide, MARGIN + Inches(0.2), y, Inches(5.6), Inches(0.56), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
        add_text(slide, MARGIN + Inches(0.36), y + Inches(0.06), Inches(1.6), Inches(0.45),
                 [[(k, 9.5, True, TEAL_DARK)]])
        add_text(slide, MARGIN + Inches(2.0), y + Inches(0.06), Inches(3.75), Inches(0.45),
                 [[(v, 8.5, False, INK)]])
        y += Inches(0.62)
    # 右上：状态机
    add_text(slide, Inches(6.85), Inches(1.5), Inches(4), Inches(0.35),
             [[("工单状态机（Leader 维护）", 12.5, True, NAVY)]])
    states = ["NEW", "TRIAGED", "PLANNED", "APPROVED", "EXECUTED", "VERIFIED", "CLOSED"]
    sx, sy = Inches(6.85), Inches(1.98)
    for i, st in enumerate(states[:4]):
        label(slide, sx + i * Inches(1.42), sy, Inches(1.26), Inches(0.42), st, LBLUE, color=NAVY, size=8.5)
        if i < 3:
            arrow(slide, sx + i * Inches(1.42) + Inches(1.26), sy + Inches(0.21),
                  sx + (i + 1) * Inches(1.42), sy + Inches(0.21), color=GRAY, weight=1.25)
    for i, st in enumerate(states[3:]):
        label(slide, sx + i * Inches(1.42), sy + Inches(0.62), Inches(1.26), Inches(0.42), st, LBLUE, color=NAVY, size=8.5)
        if i < 3:
            arrow(slide, sx + i * Inches(1.42) + Inches(1.26), sy + Inches(0.83),
                  sx + (i + 1) * Inches(1.42), sy + Inches(0.83), color=GRAY, weight=1.25)
    label(slide, Inches(6.85) + Inches(3.78), sy + Inches(1.4), Inches(1.7), Inches(0.42),
          "ESCALATED", AMBER, size=8.5)
    arrow(slide, sx + Inches(4.0), sy + Inches(0.83), sx + Inches(4.3), sy + Inches(1.4), color=AMBER, weight=1.25, dash="dash")
    # 右下：AgentTeams 映射
    add_text(slide, Inches(6.85), Inches(3.55), Inches(4), Inches(0.35),
             [[("→ AgentTeams 能力映射", 12.5, True, NAVY)]])
    maps = [
        ("Matrix 房间", "任务消息、工具结果、中间结论实时传递，Human 可见可介入"),
        ("共享存储", "对象存储 / 向量库承载 Ticket Context 与证据包"),
        ("状态追踪", "Leader 维护状态机；心跳监控 Worker 运行状态"),
        ("审计与回滚", "全部上下文变更留痕，支持审批、回滚与复盘"),
    ]
    y = Inches(4.05)
    for t, d in maps:
        rect(slide, Inches(6.85), y, Inches(5.9), Inches(0.62), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        rect(slide, Inches(6.85), y, Inches(0.07), Inches(0.62), fill=TEAL)
        add_text(slide, Inches(7.05), y + Inches(0.06), Inches(1.5), Inches(0.5),
                 [[(t, 10, True, TEAL_DARK)]])
        add_text(slide, Inches(8.55), y + Inches(0.06), Inches(4.15), Inches(0.5),
                 [[(d, 8.5, False, INK)]])
        y += Inches(0.68)
    return slide


def s09_skills(prs):
    slide = blank_slide(prs)
    header(slide, "Skill 体系设计：任务能力抽象层",
           subtitle="Skill 与 MCP 分离：Skill 抽象“做什么”，MCP 连接“怎么做”", page=9, total=TOTAL)
    layers = [
        ("沉淀层", "review.case", "复盘 Agent", AMBER),
        ("核验层", "verify.result", "核验 Agent", TEAL),
        ("执行层", "exec.refund / exec.account_change", "执行 Agent", BLUE),
        ("决策层", "plan.policy_rag", "方案 Agent", TEAL),
        ("认知层", "triage.intent", "分诊 Agent", BLUE),
        ("接入层", "intake.ingest", "聚合 Agent", NAVY),
    ]
    y = Inches(1.55)
    for name, skill, owner, color in layers:
        rect(slide, MARGIN, y, Inches(6.0), Inches(0.78), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        rect(slide, MARGIN, y, Inches(1.35), Inches(0.78), fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
        add_text(slide, MARGIN, y + Inches(0.22), Inches(1.35), Inches(0.4),
                 [[(name, 11, True, WHITE)]], align=PP_ALIGN.CENTER)
        add_text(slide, MARGIN + Inches(1.55), y + Inches(0.09), Inches(3.1), Inches(0.35),
                 [[(skill, 11, True, NAVY)]])
        add_text(slide, MARGIN + Inches(1.55), y + Inches(0.44), Inches(4.3), Inches(0.3),
                 [[("调用方：" + owner, 8.5, False, GRAY)]])
        y += Inches(0.87)
    # 右侧：九要素 + 设计要点
    add_text(slide, Inches(6.95), Inches(1.5), Inches(5.8), Inches(0.35),
             [[("每个 Skill 必须包含九要素", 12.5, True, NAVY)]])
    nine = ["① 名称　② 用途　③ 输入/输出 Schema", "④ 调用条件　⑤ 依赖工具　⑥ 失败处理机制",
            "⑦ 安全边界　⑧ 复用价值　⑨ 与多 Agent 协同流程的关系"]
    rect(slide, Inches(6.95), Inches(1.95), Inches(5.8), Inches(1.35), fill=OFFWHITE, line=LGRAY,
         line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    add_text(slide, Inches(7.2), Inches(2.08), Inches(5.35), Inches(1.1),
             [[(t, 10, False, INK)] for t in nine],
             )
    add_text(slide, Inches(6.95), Inches(3.5), Inches(5.8), Inches(0.35),
             [[("设计要点", 12.5, True, NAVY)]])
    points = [
        ("可复用", "同一 Skill 可挂载到不同 Worker，跨团队复用"),
        ("可观测", "Skill 调用全链路入 Trace，失败可回放"),
        ("可治理", "注册至 AgentTeams 统一资产库，版本化发布"),
        ("可降级", "Skill 内部封装重试 / 降级 / 人工兜底策略"),
    ]
    y = Inches(3.98)
    for t, d in points:
        rect(slide, Inches(6.95), y, Inches(5.8), Inches(0.62), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        rect(slide, Inches(6.95), y, Inches(0.07), Inches(0.62), fill=TEAL)
        add_text(slide, Inches(7.15), y + Inches(0.06), Inches(1.05), Inches(0.5),
                 [[(t, 10.5, True, TEAL_DARK)]])
        add_text(slide, Inches(8.25), y + Inches(0.06), Inches(4.45), Inches(0.5),
                 [[(d, 9, False, INK)]])
        y += Inches(0.7)
    add_text(slide, Inches(6.95), Inches(6.78), Inches(5.8), Inches(0.35),
             [[("完整九要素清单见《docs/Skill清单.md》", 9, False, GRAY)]])
    return slide


def s10_skill_sample(prs):
    slide = blank_slide(prs)
    header(slide, "核心 Skill 样例：exec.refund（退款执行）",
           subtitle="以九要素完整示例说明 Skill 的规范化设计", page=10, total=TOTAL)
    rows = [
        ["要素", "说明"],
        ["名称 / 用途", "exec.refund：方案获批后按幂等规则调用支付网关完成退款并回写工单"],
        ["输入 / 输出", "输入 {ticket_id, refund_order, amount, currency, approved_by, approval_evidence, idempotency_key} → 输出 {status, refund_txn_id, evidence, timestamps}"],
        ["调用条件", "方案已批准；风险等级 ≤ 自动执行阈值；高风险携带人工审批凭证；账户白名单校验通过"],
        ["依赖工具", "支付网关 MCP Server、工单系统 MCP Server、审计日志服务"],
        ["失败处理", "指数退避重试 ≤3 次；幂等键冲突查既有结果不重复退款；连续失败转人工并生成回滚工单"],
        ["安全边界", "白名单账户 + 额度限制；密钥由 AgentTeams 注入不进 Prompt；全程审计留痕"],
        ["复用价值", "换货、退票、改签等资金操作复用同一“幂等执行 + 证据回写 + 失败升级”模式"],
        ["协同关系", "执行 Agent 调用；上游接方案 Agent 的批准方案；下游交核验 Agent；异常经主控升级人工"],
    ]
    add_table(slide, MARGIN, Inches(1.5), Inches(7.9), rows, [1.3, 6.6], row_height=0.56, font_size=9,
              align_map={0: PP_ALIGN.CENTER})
    # 右侧：调用时序
    add_text(slide, Inches(8.7), Inches(1.5), Inches(4), Inches(0.35),
             [[("调用时序", 12.5, True, NAVY)]])
    seq = [
        ("方案 Agent", "批准方案（含审批凭证）"),
        ("执行 Agent", "调用 exec.refund"),
        ("支付网关 MCP", "退款受理 · 幂等校验"),
        ("工单系统 MCP", "回写执行证据"),
        ("核验 Agent", "比对凭证 · 满意度确认"),
    ]
    y = Inches(1.95)
    for name, d in seq:
        rect(slide, Inches(8.7), y, Inches(4.0), Inches(0.72), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        rect(slide, Inches(8.7), y, Inches(0.07), Inches(0.72), fill=TEAL)
        add_text(slide, Inches(8.92), y + Inches(0.07), Inches(1.7), Inches(0.3),
                 [[(name, 10, True, NAVY)]])
        add_text(slide, Inches(8.92), y + Inches(0.38), Inches(3.65), Inches(0.3),
                 [[(d, 8.5, False, GRAY)]])
        y += Inches(0.82)
        if name != "核验 Agent":
            arrow(slide, Inches(10.7), y - Inches(0.1), Inches(10.7), y, color=GRAY, weight=1.25)
    return slide


def s11_mcp(prs):
    slide = blank_slide(prs)
    header(slide, "MCP 与工具集成（外部系统连接层）",
           subtitle="本方案采用 MCP 协议；未落地前以等价工具契约先行，迁移成本仅协议适配层", page=11, total=TOTAL)
    rows = [
        ["工具", "调用入口", "关键参数（Schema）", "权限范围", "失败重试", "幂等控制", "审计与降级"],
        ["工单系统", "ticket.create / update",
         "{ticket_id, status, assignee, note}", "工单读写（本部门）",
         "重试 3 次 + 退避", "按 ticket_id 唯一", "审计全量；故障时降级邮件工单"],
        ["支付网关", "refund.create",
         "{order_id, amount, currency, idempotency_key}", "退款额度受限额策略约束",
         "重试 3 次；失败回滚", "幂等键全局唯一", "审计含审批凭证；降级转人工"],
        ["CRM/账户", "account.update",
         "{customer_id, field, old, new}", "白名单字段，二次确认",
         "重试 2 次", "变更前后值对拍", "审计含变更前后值；降级转人工"],
        ["知识库", "kb.search",
         "{query, top_k, filters}", "只读",
         "重试 2 次", "无状态", "审计查询；降级返回缓存"],
        ["通知渠道", "notify.send",
         "{channel, template, params}", "仅模板发送，禁自定义内容",
         "重试 3 次", "消息 ID 去重", "审计；降级换渠道"],
    ]
    add_table(slide, MARGIN, Inches(1.5), Inches(8.55), rows,
              [1.05, 1.3, 1.9, 1.5, 1.05, 1.15, 1.6], row_height=0.56, font_size=8)
    # 右侧：迁移路径
    add_text(slide, Inches(9.45), Inches(1.5), Inches(3.4), Inches(0.35),
             [[("从契约到 MCP 的迁移路径", 12, True, NAVY)]])
    path = [
        ("等价工具契约（初赛）", "REST/HTTP + OAuth2 · 参数 Schema · 审计 · 幂等", BLUE),
        ("MCP Server 适配器（复赛）", "将契约包装为 MCP 工具：tools/list · tools/call", TEAL),
        ("Agent 稳定调用", "Skill 只依赖工具名与 Schema，不感知协议差异", NAVY),
    ]
    y = Inches(1.98)
    for i, (t, d, c) in enumerate(path):
        rect(slide, Inches(9.45), y, Inches(3.35), Inches(1.15), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        rect(slide, Inches(9.45), y, Inches(0.09), Inches(1.15), fill=c)
        add_text(slide, Inches(9.7), y + Inches(0.1), Inches(2.95), Inches(0.35),
                 [[(t, 10.5, True, NAVY)]])
        add_text(slide, Inches(9.7), y + Inches(0.5), Inches(2.95), Inches(0.6),
                 [[(d, 8.5, False, GRAY)]])
        y += Inches(1.32)
        if i < 2:
            arrow(slide, Inches(11.1), y - Inches(0.17), Inches(11.1), y, color=GRAY, weight=1.5)
    add_text(slide, Inches(9.45), Inches(6.1), Inches(3.35), Inches(0.8),
             [[("迁移成本：仅协议适配，不改工具调用链、Skill 接口与安全策略。", 9, True, TEAL_DARK)]])
    return slide


def s12_observability(prs):
    slide = blank_slide(prs)
    header(slide, "可观测设计：Trace / Log / Metrics",
           subtitle="遵循 OpenTelemetry GenAI 语义，覆盖 Skill、MCP、RAG、LLM 全链路", page=12, total=TOTAL)
    cards = [
        ("Trace 全链路轨迹", ["每工单一个 TraceID 贯穿闭环", "Skill 调用 / MCP 工具 / RAG 检索 / LLM 推理", "证据包 + 执行截图可回放"],
         TEAL),
        ("Log 审计日志", ["审批记录、权限变更、异常栈", "敏感操作留痕，满足合规审计", "与 TraceID 关联检索"],
         BLUE),
        ("Metrics 业务与成本指标", ["工单时效 / 一次解决率 / 人工介入率", "Token 消耗 / 模型调用量 / 错误率", "告警：异常率、超时、额度超限"],
         AMBER),
    ]
    x = MARGIN
    for t, lines, c in cards:
        rect(slide, x, Inches(1.5), Inches(3.94), Inches(2.0), fill=WHITE, line=LGRAY,
             line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
        rect(slide, x, Inches(1.5), Inches(3.94), Inches(0.46), fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
        add_text(slide, x, Inches(1.56), Inches(3.94), Inches(0.36),
                 [[(t, 12.5, True, WHITE)]], align=PP_ALIGN.CENTER)
        add_text(slide, x + Inches(0.2), Inches(2.1), Inches(3.55), Inches(1.3),
                 [[("• " + ln, 9.5, False, INK)] for ln in lines],
                 )
        x += Inches(4.08)
    # 采集与消费链路
    add_text(slide, MARGIN, Inches(3.75), Inches(6), Inches(0.35),
             [[("采集与消费链路", 12.5, True, NAVY)]])
    flow = ["Agent / Skill / MCP / RAG / LLM 推理", "OTLP 采集器", "SLS 存储 + AgentLoop 观测", "实时监控告警 / 离线评估"]
    x = MARGIN
    for i, f in enumerate(flow):
        w = Inches(text_width_in(f, 9.5))
        label(slide, x, Inches(4.2), w, Inches(0.55), f, OFFWHITE, color=NAVY, size=9.5, line=LGRAY)
        x += w + Inches(0.12)
        if i < len(flow) - 1:
            arrow(slide, x - Inches(0.12), Inches(4.47), x, Inches(4.47), color=TEAL, weight=1.5)
    # 应用场景
    apps = [
        ("在线监控", "仪表盘实时查看 Worker 用量、模型调用、Token 趋势，异常自动告警"),
        ("离线评估", "基于评估集量化意图识别准确率、方案合规率、一次解决率，驱动优化"),
        ("评审证据", "每工单输出 Trace 报告 + 证据包（截图/日志/结果凭证）作为执行证据沉淀"),
    ]
    y = Inches(5.1)
    for t, d in apps:
        rect(slide, MARGIN, y, W - 2 * MARGIN, Inches(0.58), fill=OFFWHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        rect(slide, MARGIN, y, Inches(0.07), Inches(0.58), fill=TEAL)
        add_text(slide, MARGIN + Inches(0.25), y + Inches(0.06), Inches(1.5), Inches(0.45),
                 [[(t, 11, True, TEAL_DARK)]])
        add_text(slide, MARGIN + Inches(1.85), y + Inches(0.06), W - 2 * MARGIN - Inches(2.1), Inches(0.45),
                 [[(d, 9.5, False, INK)]])
        y += Inches(0.68)
    return slide


def s13_rag(prs):
    slide = blank_slide(prs)
    header(slide, "RAG 与上下文增强",
           subtitle="四项上下文能力全覆盖，满足赛题“至少实现 2 项”；RAG 贯穿检索 → 证据对齐 → 决策", page=13, total=TOTAL)
    caps = [
        ("Agent 记忆存储", "短期：工单上下文（Redis）\n长期：历史案例向量化", TEAL),
        ("知识库 RAG", "SOP / 产品手册 / 历史工单\n混合检索 + 重排 + 引用溯源", BLUE),
        ("共享状态管理", "Ticket Context 单一事实源\n状态机驱动流转", AMBER),
        ("轨迹可观测", "TraceID 贯穿全链路\n证据包可回放可审计", NAVY),
    ]
    pos = [(MARGIN, Inches(1.55)), (MARGIN + Inches(3.12), Inches(1.55)),
           (MARGIN, Inches(3.15)), (MARGIN + Inches(3.12), Inches(3.15))]
    for (x, y), (t, d, c) in zip(pos, caps):
        rect(slide, x, y, Inches(2.98), Inches(1.44), fill=WHITE, line=LGRAY,
             line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.09)
        rect(slide, x, y, Inches(2.98), Inches(0.4), fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
        add_text(slide, x + Inches(0.1), y + Inches(0.04), Inches(2.8), Inches(0.34),
                 [[(t, 11, True, WHITE)]], align=PP_ALIGN.CENTER)
        lines = d.split("\n")
        add_text(slide, x + Inches(0.14), y + Inches(0.52), Inches(2.7), Inches(0.9),
                 [[("• " + ln, 8.5, False, INK)] for ln in lines],
                 )
    label(slide, MARGIN + Inches(3.15), Inches(4.72), Inches(2.9), Inches(0.4),
          "4 项全覆盖 ✓ 满足 ≥2 项要求", GREEN, size=9.5)
    # 右侧：检索增强决策流
    add_text(slide, Inches(6.7), Inches(1.5), Inches(5.6), Inches(0.35),
             [[("RAG 决策链路（防幻觉设计）", 12.5, True, NAVY)]])
    flow = [
        ("工单 + 分级", "方案 Agent 判断是否需要检索", NAVY),
        ("MCP 接入数据源", "知识库 / 历史案例 / 业务数据", BLUE),
        ("Skill 检索 + 证据对齐", "plan.policy_rag：混合检索 + 重排 + 引用", TEAL),
        ("证据充分性判断", "Agent 判定证据是否足以支撑决策", AMBER),
        ("生成方案 / 升级人工", "证据不足则标注并升级，杜绝幻觉执行", RED),
    ]
    y = Inches(1.95)
    for t, d, c in flow:
        rect(slide, Inches(6.7), y, Inches(6.0), Inches(0.72), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        rect(slide, Inches(6.7), y, Inches(0.09), Inches(0.72), fill=c)
        add_text(slide, Inches(6.98), y + Inches(0.07), Inches(2.6), Inches(0.6),
                 [[(t, 10.5, True, NAVY)]], anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, Inches(9.6), y + Inches(0.07), Inches(3.0), Inches(0.6),
                 [[(d, 8.5, False, GRAY)]], anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.82)
        if t != "生成方案 / 升级人工":
            arrow(slide, Inches(9.7), y - Inches(0.1), Inches(9.7), y, color=GRAY, weight=1.25)
    return slide


def s14_security(prs):
    slide = blank_slide(prs)
    header(slide, "安全边界、审批与回滚、审计",
           subtitle="高风险动作人工确认：审批凭证链 + 幂等执行 + 回滚审计", page=14, total=TOTAL)
    rows = [
        ["风险级别", "典型动作", "执行策略"],
        ["低风险", "修改工单备注、发送通知、资料补全", "规则校验后自动执行，全程留痕"],
        ["中风险", "换货、非敏感信息变更", "方案 Agent 生成方案 + 规则校验，限额内自动执行"],
        ["高风险", "退款、账户敏感变更、批量操作", "人工审批（大额双人复核）后执行，支持回滚"],
    ]
    add_table(slide, MARGIN, Inches(1.5), W - 2 * MARGIN, rows,
              [1.6, 4.4, 6.23], row_height=0.5, font_size=10,
              align_map={0: PP_ALIGN.CENTER})
    # 审批与回滚
    add_text(slide, MARGIN, Inches(3.4), Inches(5), Inches(0.35),
             [[("审批与回滚机制", 12.5, True, NAVY)]])
    items1 = [
        ("审批凭证链", "审批意见、审批人、时间、动作哈希进入执行上下文"),
        ("幂等控制", "幂等键全局唯一，重复请求不重复执行"),
        ("回滚机制", "每类动作预置回滚脚本，执行失败自动回滚并通知"),
        ("额度控制", "单笔 / 单日限额，超限强制人工审批"),
    ]
    y = Inches(3.85)
    for t, d in items1:
        rect(slide, MARGIN, y, Inches(6.0), Inches(0.6), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        rect(slide, MARGIN, y, Inches(0.07), Inches(0.6), fill=AMBER)
        add_text(slide, MARGIN + Inches(0.22), y + Inches(0.05), Inches(1.4), Inches(0.5),
                 [[(t, 10, True, NAVY)]])
        add_text(slide, MARGIN + Inches(1.72), y + Inches(0.05), Inches(4.25), Inches(0.5),
                 [[(d, 8.5, False, INK)]])
        y += Inches(0.68)
    # 审计与合规
    add_text(slide, Inches(7.0), Inches(3.4), Inches(5), Inches(0.35),
             [[("审计与合规", 12.5, True, NAVY)]])
    items2 = [
        ("RBAC 最小权限", "Worker 仅持本职 Token，跨权限须主控路由"),
        ("敏感数据脱敏", "会话原文先脱敏再入库，满意度数据匿名化"),
        ("密钥治理", "密钥由 AgentTeams 注入，不进入 Prompt / 日志"),
        ("全程审计", "审批、执行、回滚全部留痕，支持合规追溯"),
    ]
    y = Inches(3.85)
    for t, d in items2:
        rect(slide, Inches(7.0), y, Inches(5.8), Inches(0.6), fill=WHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        rect(slide, Inches(7.0), y, Inches(0.07), Inches(0.6), fill=BLUE)
        add_text(slide, Inches(7.22), y + Inches(0.05), Inches(1.6), Inches(0.5),
                 [[(t, 10, True, NAVY)]])
        add_text(slide, Inches(8.9), y + Inches(0.05), Inches(3.85), Inches(0.5),
                 [[(d, 8.5, False, INK)]])
        y += Inches(0.68)
    add_text(slide, MARGIN, Inches(6.6), W - 2 * MARGIN, Inches(0.35),
             [[("安全设计贯穿 AgentTeams：团队权限隔离、模型配置管控、监控仪表盘异常识别",
                 10, False, GRAY)]])
    return slide


def s15_exception(prs):
    slide = blank_slide(prs)
    header(slide, "异常分支与容错设计",
           subtitle="每个异常都有“检测方式 + 处理策略”，确保闭环不中断、不失控", page=15, total=TOTAL)
    rows = [
        ["异常场景", "检测方式", "处理策略"],
        ["意图置信度低", "分诊 Agent 输出置信度分数", "追问补全信息；仍不满足则转人工"],
        ["重复工单", "会话指纹 / 语义相似度去重", "合并至既有工单，保留关联记录"],
        ["工具调用失败", "MCP 返回错误码 + Trace 异常", "指数退避重试 ≤3 次；失败升级人工"],
        ["执行结果异常", "核验 Agent 比对金额 / 状态 / 凭证", "标记异常 → 自动回滚或人工介入"],
        ["检索证据不足", "plan.policy_rag 返回证据得分", "方案标注“证据不足”，不执行并升级"],
        ["审批超时", "审批台超时定时器", "升级提醒；紧急工单提升优先级"],
        ["模型不可用", "健康检查 / 调用错误率", "切换备用模型或规则引擎降级"],
        ["敏感信息泄露风险", "内容合规检测 / 脱敏过滤器", "拒绝写入、告警并隔离该工单"],
    ]
    add_table(slide, MARGIN, Inches(1.5), W - 2 * MARGIN, rows,
              [3.0, 4.4, 4.83], row_height=0.56, font_size=9.5)
    add_text(slide, MARGIN, Inches(6.55), W - 2 * MARGIN, Inches(0.4),
             [[("降级原则：先保可用 → 再保正确 → 最后保自动化；任何无法自动恢复的路径都收敛到人工（Human-in-the-loop）。",
                 10.5, True, TEAL_DARK)]])
    return slide


def s16_roadmap(prs):
    slide = blank_slide(prs)
    header(slide, "可行性与落地计划",
           subtitle="四阶段推进：初赛设计 → 复赛最小闭环 → 决赛生产化验证 → 赛后开源", page=16, total=TOTAL)
    phases = [
        ("初赛 · 当前", "方向与方案设计", TEAL, [
            "方案设计 + 架构图",
            "7 个 Agent Identity",
            "7 个 Skill 九要素清单",
            "MCP 等价契约与迁移路径",
        ], "已完成"),
        ("复赛", "可运行最小闭环", BLUE, [
            "AgentTeams 本地部署",
            "3+ Worker 跑通闭环",
            "Skill 落地与样例输入输出",
            "运行证据（Trace/日志）",
        ], "进行中"),
        ("决赛", "生产化验证", AMBER, [
            "MCP 接入真实/仿真企业系统",
            "可观测 + 离线评估体系",
            "审批 / 回滚 / 安全演练",
            "压测与多租户隔离",
        ], "计划"),
        ("赛后", "开放与复用", NAVY, [
            "开源 Skill 库与闭环模板",
            "MCP 工具契约标准化",
            "复盘数据集模板",
            "跨场景复制（IT/政务/金融）",
        ], "计划"),
    ]
    x = MARGIN
    for i, (phase, tag, c, items, status) in enumerate(phases):
        rect(slide, x, Inches(1.55), Inches(2.92), Inches(3.9), fill=WHITE, line=LGRAY,
             line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        rect(slide, x, Inches(1.55), Inches(2.92), Inches(0.62), fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
        add_text(slide, x + Inches(0.15), Inches(1.62), Inches(2.6), Inches(0.3),
                 [[(phase, 12.5, True, WHITE)]])
        add_text(slide, x + Inches(0.15), Inches(2.02), Inches(2.6), Inches(0.3),
                 [[(tag, 10.5, True, WHITE)]])
        y = Inches(2.4)
        for it in items:
            add_text(slide, x + Inches(0.22), y, Inches(2.55), Inches(0.55),
                     [[("• " + it, 9.5, False, INK)]])
            y += Inches(0.52)
        label(slide, x + Inches(0.8), Inches(4.75), Inches(1.35), Inches(0.4), status,
              GREEN if status == "已完成" else (TEAL if status == "进行中" else LGRAY),
              color=WHITE if status != "计划" else GRAY, size=9.5)
        x += Inches(3.06)
    # 可行性说明
    rect(slide, MARGIN, Inches(5.7), W - 2 * MARGIN, Inches(1.0), fill=OFFWHITE, line=LGRAY,
         line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    add_text(slide, MARGIN + Inches(0.2), Inches(5.8), Inches(2.2), Inches(0.4),
             [[("可行性要点", 12, True, NAVY)]])
    feats = [
        ("组件可得", "AgentTeams 云服务 / 开源部署 + 阿里云生态，零自研框架成本"),
        ("依赖可控", "仅需 Docker + LLM API Key，2-3 人即可推进复赛闭环"),
        ("演进平滑", "契约先行 → MCP 适配 → 生产化，调用链无需重构"),
    ]
    y = Inches(5.9)
    for t, d in feats:
        add_text(slide, MARGIN + Inches(2.35), y, Inches(1.15), Inches(0.3),
                 [[(t, 10, True, TEAL_DARK)]])
        add_text(slide, MARGIN + Inches(3.5), y, Inches(9.0), Inches(0.3),
                 [[(d, 9.5, False, INK)]])
        y += Inches(0.3)
    return slide


def s17_opensource(prs):
    slide = blank_slide(prs)
    header(slide, "开放 / 开源价值",
           subtitle="输出可复用资产：闭环模板、Skill 库、MCP 契约、评估数据集", page=17, total=TOTAL)
    assets = [
        ("客服闭环编排模板", "Agent Teams 定义 + 工单状态机 + 上下文 Schema，开箱即用", TEAL),
        ("Skill 库", "7 个核心 Skill 九要素文档 + 可运行实现（复赛）", BLUE),
        ("MCP 工具契约", "工单 / 支付 / CRM / 知识库 / 通知的接口契约与适配器", AMBER),
        ("复盘数据集模板", "疑难案例标注、评估指标与复盘报告模板", NAVY),
    ]
    x = MARGIN
    for t, d, c in assets:
        rect(slide, x, Inches(1.55), Inches(2.92), Inches(1.9), fill=WHITE, line=LGRAY,
             line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
        rect(slide, x, Inches(1.55), Inches(2.92), Inches(0.08), fill=c)
        add_text(slide, x + Inches(0.16), Inches(1.78), Inches(2.6), Inches(0.7),
                 [[(t, 11.5, True, NAVY)]])
        add_text(slide, x + Inches(0.16), Inches(2.5), Inches(2.6), Inches(0.85),
                 [[(d, 8.5, False, GRAY)]])
        x += Inches(3.06)
    # 复用场景
    add_text(slide, MARGIN, Inches(3.75), Inches(6), Inches(0.35),
             [[("可复制场景", 12.5, True, NAVY)]])
    scens = ["电商售后", "IT 服务台", "政务热线", "金融客服", "物流 / 出行", "SaaS 订阅"]
    x = MARGIN
    for s in scens:
        w = Inches(0.35 + 0.22 * len(s))
        label(slide, x, Inches(4.2), w, Inches(0.5), s, LBLUE, color=NAVY, size=10.5, line=LGRAY)
        x += w + Inches(0.16)
    # 生态协同
    add_text(slide, MARGIN, Inches(4.95), Inches(6), Inches(0.35),
             [[("生态协同", 12.5, True, NAVY)]])
    eco = [
        ("AgentTeams / HiClaw 社区", "沉淀客服团队模板与 Worker 配置，反向贡献社区"),
        ("AgentLoop 观测生态", "共享评估指标与 Trace 语义，降低多团队适配成本"),
        ("企业 IM 入口", "钉钉 / 飞书 / 企业微信发起任务，天然支持 Human-in-the-loop"),
    ]
    y = Inches(5.4)
    for t, d in eco:
        rect(slide, MARGIN, y, W - 2 * MARGIN, Inches(0.55), fill=OFFWHITE, line=LGRAY,
             line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        rect(slide, MARGIN, y, Inches(0.07), Inches(0.55), fill=TEAL)
        add_text(slide, MARGIN + Inches(0.25), y + Inches(0.05), Inches(3.6), Inches(0.45),
                 [[(t, 10.5, True, NAVY)]])
        add_text(slide, MARGIN + Inches(4.0), y + Inches(0.05), W - 2 * MARGIN - Inches(4.25), Inches(0.45),
                 [[(d, 9, False, INK)]])
        y += Inches(0.63)
    return slide


def s18_progress(prs):
    slide = blank_slide(prs)
    header(slide, "项目进展与总结", page=18, total=TOTAL)
    # 已完成
    rect(slide, MARGIN, Inches(1.55), Inches(6.0), Inches(3.6), fill=OFFWHITE, line=LGRAY,
         line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(slide, MARGIN, Inches(1.55), Inches(6.0), Inches(0.5), fill=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
    add_text(slide, MARGIN + Inches(0.2), Inches(1.61), Inches(5.6), Inches(0.4),
             [[("已完成（初赛）", 13, True, WHITE)]])
    done = [
        "方案设计：Leader-Worker 五层架构 + 八环节闭环",
        "7 个 Agent 的 Identity 清单（附录 A）",
        "7 个核心 Skill 九要素完整定义",
        "MCP 工具契约与迁移路径设计",
        "样例输入输出与校验脚本（可运行）",
        "本方案 PPT 与交付文档",
    ]
    y = Inches(2.2)
    for d in done:
        add_text(slide, MARGIN + Inches(0.3), y, Inches(5.5), Inches(0.5),
                 [[("✓ ", 10.5, True, GREEN), (d, 10, False, INK)]])
        y += Inches(0.5)
    # 下一步
    rect(slide, Inches(6.9), Inches(1.55), Inches(5.9), Inches(3.6), fill=OFFWHITE, line=LGRAY,
         line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(slide, Inches(6.9), Inches(1.55), Inches(5.9), Inches(0.5), fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
    add_text(slide, Inches(7.1), Inches(1.61), Inches(5.4), Inches(0.4),
             [[("下一步（复赛）", 13, True, WHITE)]])
    todo = [
        "基于 AgentTeams 本地部署，创建 3+ Worker",
        "跑通最小闭环：聚合 → 分诊 → 方案 → 执行 → 核验",
        "落地 2-3 个 Skill 与 MCP 适配器",
        "产出样例输入输出与运行证据（Trace / 日志 / 截图）",
        "接入离线评估，量化一次解决率与人工介入率",
    ]
    y = Inches(2.2)
    for d in todo:
        add_text(slide, Inches(7.2), y, Inches(5.3), Inches(0.5),
                 [[("→ ", 10.5, True, BLUE), (d, 10, False, INK)]])
        y += Inches(0.5)
    # 总结语
    rect(slide, MARGIN, Inches(5.5), W - 2 * MARGIN, Inches(1.2), fill=NAVY,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    add_text(slide, MARGIN + Inches(0.4), Inches(5.72), W - 2 * MARGIN - Inches(0.8), Inches(0.8),
             [[("一句话总结：", 13, True, CYAN),
               ("让客服从“单点问答”走向“自主闭环”——多 Agent 协同完成接入、分诊、方案、执行、核验与沉淀，",
                12, False, WHITE)],
              [("可执行、可核验、可复盘、可进化，并以 AgentTeams 为底座平滑走向生产。", 12, True, WHITE)]])
    return slide


def s19_end(prs):
    slide = blank_slide(prs, bg=NAVY_DEEP)
    circle(slide, Inches(11.9), Inches(6.35), Inches(2.2), line=RGBColor(0x1B, 0x3A, 0x5F), line_w=1.5)
    circle(slide, Inches(1.1), Inches(1.0), Inches(0.7), fill=TEAL)
    add_text(slide, 0, Inches(2.7), W, Inches(1.0),
             [[("谢谢观看", 44, True, WHITE)]], align=PP_ALIGN.CENTER)
    add_text(slide, 0, Inches(3.85), W, Inches(0.5),
             [[("智服方舟 · 多智能体客服自主闭环平台", 18, True, CYAN)]], align=PP_ALIGN.CENTER)
    add_text(slide, 0, Inches(4.55), W, Inches(0.4),
             [[("恳请各位评委老师指正", 13, False, GRAY_LIGHT)]], align=PP_ALIGN.CENTER)
    add_text(slide, 0, Inches(6.4), W, Inches(0.35),
             [[("团队成员：____________　　指导教师：____________　　联系方式：____________", 10, False, GRAY_LIGHT)]],
             align=PP_ALIGN.CENTER)
    return slide


def main():
    prs = new_prs()
    for fn in (s01_cover, s02_toc, s03_problem, s04_arch, s05_roles, s06_identity,
               s07_loop, s08_context, s09_skills, s10_skill_sample, s11_mcp,
               s12_observability, s13_rag, s14_security, s15_exception, s16_roadmap,
               s17_opensource, s18_progress, s19_end):
        fn(prs)
    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "ppt")
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "智服方舟-初赛方案.pptx")
    prs.save(out)
    print("PPT saved:", out, "| slides:", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
